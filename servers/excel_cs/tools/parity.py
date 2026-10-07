"""Bounded Python/C# cell-edit parity scenario, checked against OOXML directly."""

import asyncio
import posixpath
import subprocess
import sys
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from urllib.parse import unquote

from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client


ROOT = Path(__file__).resolve().parents[3]
EXCEL = ROOT / "servers" / "excel"
SERVER = ROOT / "servers" / "excel_cs"
MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE = "http://schemas.openxmlformats.org/package/2006/relationships"
VALUES = {"B1": 27, "C1": 9, "E5": True,
          "F6": 0.1, "G7": 4, "H7": 4, "I8": 12, "J8": "batch"}
EXPECTED = {**{address: (float(value) if type(value) in (int, float) else value, None)
               for address, value in VALUES.items()}, "D3": (None, None), "K9": (5.0, "2+3")}
VARIANTS = ("default", "prefixed-x", "bom-crlf-standalone", "opc-percent-case",
            "new-shared-strings", "nested-workbook")


def dotnet(*arguments):
    subprocess.run(["dotnet", *arguments], cwd=ROOT, check=True, stdout=subprocess.DEVNULL)


def read_cells(path, include_facets=False):
    """Resolve OPC relationships and read cells without either implementation's locator."""
    with zipfile.ZipFile(path) as archive:
        entries = {name.lower(): name for name in archive.namelist()}

        def part(name):
            return archive.read(entries[name.lower()])

        def relationships(name):
            base = "" if name == "_rels/.rels" else posixpath.dirname(posixpath.dirname(name))
            return {node.attrib["Id"]: (node.attrib["Type"].split("/")[-1],
                    posixpath.normpath(posixpath.join(base, unquote(node.attrib["Target"]).lstrip("/")))
                    if not node.attrib["Target"].startswith("/") else unquote(node.attrib["Target"]).lstrip("/"))
                    for node in ET.fromstring(part(name)).findall(f"{{{PACKAGE}}}Relationship")
                    if node.attrib.get("TargetMode") != "External"}

        workbook = next(target for kind, target in relationships("_rels/.rels").values()
                        if kind == "officeDocument")
        book = ET.fromstring(part(workbook))
        rels = relationships(posixpath.join(posixpath.dirname(workbook), "_rels",
                                          posixpath.basename(workbook) + ".rels"))
        sheet = next(node for node in book.iter(f"{{{MAIN}}}sheet") if node.attrib["name"] == "Sheet1")
        worksheet = ET.fromstring(part(rels[sheet.attrib[f"{{{OFFICE}}}id"]][1]))
        strings = []
        shared_items = []
        for kind, target in rels.values():
            if kind == "sharedStrings":
                for item in ET.fromstring(part(target)).findall(f"{{{MAIN}}}si"):
                    shared_items.append(item)
                    strings.append("".join(node.text or "" for node in
                                           item.findall(f"{{{MAIN}}}t") + item.findall(f"{{{MAIN}}}r/{{{MAIN}}}t")))
        result = {}
        cells = {}
        for cell in worksheet.iter(f"{{{MAIN}}}c"):
            cells[cell.attrib["r"]] = cell
            formula = cell.find(f"{{{MAIN}}}f")
            value = cell.find(f"{{{MAIN}}}v")
            kind = cell.attrib.get("t", "n")
            if kind == "s":
                scalar = strings[int(value.text)] if value is not None and value.text else None
            elif kind == "inlineStr":
                scalar = "".join(node.text or "" for node in cell.findall(f"{{{MAIN}}}is//{{{MAIN}}}t"))
            elif kind == "b":
                scalar = value is not None and value.text == "1"
            elif kind == "n" and value is not None and value.text is not None:
                scalar = float(value.text)
            else:
                scalar = value.text if value is not None else None
            result[cell.attrib["r"]] = (scalar, formula.text if formula is not None else None)
        if not include_facets:
            return result

        def shape(element):
            if element is None:
                return None
            attributes = dict(element.attrib)
            if element.tag == f"{{{MAIN}}}b" and attributes.get("val") in ("1", "true"):
                del attributes["val"]
            tag = f"{{{MAIN}}}si" if element.tag == f"{{{MAIN}}}is" else element.tag
            return (tag, tuple(sorted(attributes.items())), element.text or "",
                    tuple(shape(child) for child in element))

        rich_cell = cells.get("A1")
        rich_item = None
        if rich_cell is not None:
            if rich_cell.attrib.get("t") == "s":
                rich_item = shared_items[int(rich_cell.find(f"{{{MAIN}}}v").text)]
            else:
                rich_item = rich_cell.find(f"{{{MAIN}}}is")
        facets = {"A1.rich": shape(rich_item)}
        style_part = next(target for kind, target in rels.values() if kind == "styles")
        styles = ET.fromstring(part(style_part))
        xfs = styles.find(f"{{{MAIN}}}cellXfs")
        for address in ("B1", "D3"):
            if address not in cells:
                facets[f"{address}.style"] = None
                continue
            style_id = int(cells[address].attrib.get("s", "0"))
            xf = xfs[style_id]
            parts = []
            for attr, group in (("fontId", "fonts"), ("fillId", "fills"), ("borderId", "borders")):
                index = int(xf.attrib.get(attr, "0"))
                parts.append(shape(styles.find(f"{{{MAIN}}}{group}")[index]))
            facets[f"{address}.style"] = (shape(xf), *parts)
        return result, facets


def check_output(source, output, include_facets=False):
    if include_facets:
        original, original_facets = read_cells(source, True)
        actual, actual_facets = read_cells(output, True)
    else:
        original = read_cells(source)
        actual = read_cells(output)
    errors = []
    for address, expected in EXPECTED.items():
        if address not in actual or actual[address] != expected:
            errors.append(f"{address}: expected {expected!r}, got {actual.get(address)!r}")
    for address in original.keys() - EXPECTED.keys():
        if original[address] != actual.get(address):
            errors.append(f"{address}: unrelated cell changed: {original[address]!r} -> {actual.get(address)!r}")
    if include_facets:
        for name, expected in original_facets.items():
            if actual_facets[name] != expected:
                errors.append(f"{name}: preserved facet changed")
    return errors


def create_independent_fixture(path):
    from openpyxl import Workbook
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet["A1"] = CellRichText("plain", TextBlock(InlineFont(b=True, color="FFFF0000"), " bold"))
    sheet["B1"] = 42
    sheet["B1"].font = Font(name="Arial", bold=True, color="FF0070C0")
    sheet["B1"].fill = PatternFill("solid", fgColor="FFFFE699")
    sheet["B1"].number_format = "0.00"
    sheet["C1"] = "=1+1"
    sheet["D3"] = "old"
    sheet["D3"].font = Font(italic=True, color="FF008000")
    workbook.save(path)


def check_facet_negative_controls(source, output, corrupted):
    with zipfile.ZipFile(output) as original, zipfile.ZipFile(corrupted, "w") as altered:
        for entry in original.infolist():
            content = original.read(entry)
            if entry.filename.lower() == "xl/worksheets/sheet1.xml":
                worksheet = ET.fromstring(content)
                cells = {cell.attrib["r"]: cell for cell in worksheet.iter(f"{{{MAIN}}}c")}
                rich_run = cells["A1"].findall(f"{{{MAIN}}}is/{{{MAIN}}}r")[-1]
                rich_run.remove(rich_run.find(f"{{{MAIN}}}rPr"))
                cells["B1"].attrib["s"] = "0"
                content = ET.tostring(worksheet, encoding="utf-8")
            altered.writestr(entry, content)
    expected = ["A1.rich: preserved facet changed", "B1.style: preserved facet changed"]
    actual = check_output(source, corrupted, True)
    if actual != expected:
        raise AssertionError(f"Facet negative controls: expected {expected}, got {actual}")


async def save_new(source, output, dll, assert_rich=False):
    parameters = StdioServerParameters(command="dotnet", args=[str(dll)], cwd=str(ROOT))
    async with stdio_client(parameters) as (reader, writer):
        async with ClientSession(reader, writer) as client:
            await client.initialize()

            async def call(name, arguments):
                response = await client.call_tool(name, arguments)
                data = response.structuredContent
                if response.isError or data is None or not data.get("ok"):
                    raise RuntimeError(f"{name}: {data}")
                return data["data"]

            opened = await call("excel_open", {"path": str(source)})
            session = opened["session"]
            await call("excel_apply", {"session": session, "base_revision": 0, "sheet": "Sheet1", "ops": [
                *({"op": "set_value", "target": address, "value": VALUES[address]}
                  for address in ("B1", "C1", "E5", "F6")),
                {"op": "clear", "target": "D3", "what": ["values"]},
                {"op": "set_formula", "target": "K9", "formula": "=2+3", "cache": {"value": 5}},
                {"op": "fill", "target": "G7:H7", "value": 4},
                {"op": "set_values", "target": "I8:J8", "values": [[12, "batch"]]},
            ]})
            save = {"session": session, "mode": "copy", "path": str(output)}
            if assert_rich:
                save["assert"] = [{"target": "Sheet1!A1", "equals": {
                    "value": "plain bold", "rich": '<r>plain</r><r b color="FF0000"> bold</r>'}}]
            await call("excel_save", save)
            await call("excel_close", {"session": session, "discard_unsaved": True})


def save_legacy(source, output):
    sys.path.insert(0, str(EXCEL))
    import main

    main.excel_load(str(source))
    session = str(source.resolve())
    main.excel_edit_cells(session, "Sheet1", [{"cell": address, "value": value}
                                              for address, value in VALUES.items()])
    main.excel_clear_range(session, "Sheet1", 2, 3, 2, 3)
    main.excel_set_formula(session, "Sheet1", "K9", "=2+3", cached_value=5, cached_value_present=True)
    main.excel_save_as_copy(session, str(output), report_format="json", verify_preservation=False)
    main.excel_close(session)


async def main():
    dotnet("build", str(SERVER / "src/DocLoupe.Excel.Server/DocLoupe.Excel.Server.csproj"), "-c", "Release", "--nologo")
    dll = SERVER / "src/DocLoupe.Excel.Server/bin/Release/net10.0/DocLoupe.Excel.Server.dll"
    with tempfile.TemporaryDirectory(prefix="docloupe-parity-") as temporary:
        directory = Path(temporary)
        dotnet("run", "--project", str(SERVER / "fixtures/DocLoupe.Excel.Fixtures.csproj"),
               "-c", "Release", "--", str(directory))
        create_independent_fixture(directory / "openpyxl-rich-style.xlsx")
        failures = []
        for variant in (*VARIANTS, "openpyxl-rich-style"):
            source = directory / f"{variant}.xlsx"
            current = directory / f"{variant}-cs.xlsx"
            legacy = directory / f"{variant}-python.xlsx"
            include_facets = variant == "openpyxl-rich-style"
            await save_new(source, current, dll, include_facets)
            current_errors = check_output(source, current, include_facets)
            if current_errors:
                failures.append(f"{variant} C#: {current_errors}")
            if include_facets:
                check_facet_negative_controls(source, current, directory / "corrupted-facets.xlsx")
            try:
                save_legacy(source, legacy)
                legacy_errors = check_output(source, legacy, include_facets)
            except (KeyError, ValueError) as error:
                legacy_errors = [f"{type(error).__name__}: {error}"]
            if variant == "opc-percent-case":
                expected_errors = ["ValueError: Sheet 'Sheet1' not found. Available: []"]
            elif variant in ("new-shared-strings", "openpyxl-rich-style"):
                expected_errors = []
            elif variant in ("prefixed-x", "nested-workbook"):
                expected_errors = ["D3: expected (None, None), got None",
                                   "A1: unrelated cell changed: ('hello', None) -> None"]
            else:
                expected_errors = ["A1: unrelated cell changed: ('hello', None) -> ('hellohe', None)"]
            if legacy_errors != expected_errors:
                failures.append(f"{variant} Python: expected {expected_errors}, got {legacy_errors}")
            print(f"{variant}: C# {'PASS' if not current_errors else 'FAIL'}; "
                  f"Python {'PASS' if not legacy_errors else 'known divergence: ' + str(legacy_errors)}")
        if failures:
            raise AssertionError("\n".join(failures))
    print("Parity scenario PASS: 7 C# outputs, 2 comparable Python outputs, 5 recorded divergences")


if __name__ == "__main__":
    asyncio.run(main())
