import json
import hashlib
import shutil
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "servers" / "excel"))

import main as M
from core import inspect_xlsx_package
from save_transaction import _requested_package_additions


def _session_key(result):
    return result.split("session_key=", 1)[1].split(" |", 1)[0].strip("'")


def _rewrite_part(source, output, name, old, new):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as modified:
        for info in original.infolist():
            raw = original.read(info.filename)
            if info.filename == name:
                raw = raw.replace(old, new)
            modified.writestr(info, raw)


def test_undeclared_markup_compatibility_prefix_is_invalid_and_not_equivalent(tmp_path):
    source = tmp_path / "source.xlsx"
    changed = tmp_path / "changed.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "data"
    workbook.save(source)
    with zipfile.ZipFile(source) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
    _rewrite_part(source, changed, "xl/worksheets/sheet1.xml", sheet.split(b">", 1)[0] + b">",
                  sheet.split(b">", 1)[0] + b' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="unknown">')
    report = inspect_xlsx_package(str(changed))
    assert report["valid"] is False
    assert "undeclared mc:Ignorable" in str(report["errors"])
    with pytest.raises(ValueError, match="Invalid markup compatibility"):
        M.excel_verify_preservation(str(changed), before_path=str(source))


def test_bulk_cell_edit_error_keeps_entire_session_unchanged(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet"
    workbook.active["B3"] = "original"
    workbook.active["D1"] = "rich"
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    data = M._get_session(key)
    sheet = M._find_sheet(data, "Sheet")
    sheet["rows"][0]["cells"][3]["rich_text"] = {"runs": [{"text": "rich"}]}
    before = json.loads(M.excel_get_cell(key, "Sheet", 2, 1))["value"]
    dirty_paths = list(data.get("_dirty_paths") or [])
    with pytest.raises(ValueError, match="rich.text"):
        M.excel_edit_cells(key, "Sheet", [{"cell": "B3", "value": "changed"},
                                           {"cell": "D1", "value": "plain"}])
    assert json.loads(M.excel_get_cell(key, "Sheet", 2, 1))["value"] == before
    assert list(data.get("_dirty_paths") or []) == dirty_paths


def test_read_only_tools_and_reload_preserve_source(tmp_path):
    source = tmp_path / "read-only.xlsx"
    picture = tmp_path / "picture.png"
    Image.new("RGB", (2, 2), (20, 30, 40)).save(picture)
    workbook = openpyxl.Workbook()
    workbook.active.title = "Sheet"
    workbook.active["A1"] = "Header"
    workbook.active["A2"] = "alpha"
    workbook.active.add_image(openpyxl.drawing.image.Image(str(picture)), "B2")
    workbook.save(source)
    original_hash = hashlib.sha256(source.read_bytes()).hexdigest()

    session_key = _session_key(M.excel_load(str(source)))
    try:
        assert json.loads(M.excel_find_rows(session_key, "Sheet", 0, value="alpha"))[0]["row_index"] == 1
        assert "alpha" in M.excel_to_markdown(session_key, "Sheet").text
        result = M.excel_extract_images(session_key, "Sheet", str(tmp_path / "images"))
        assert "image_01.png" in result
        assert (tmp_path / "images" / "image_01.png").is_file()

        M.excel_edit_cells(session_key, "Sheet", [{"cell": "A2", "value": "temporary"}])
        assert json.loads(M.excel_find_rows(session_key, "Sheet", 0, value="temporary"))
        assert "Reloaded:" in M.excel_reload(session_key)
        assert json.loads(M.excel_find_rows(session_key, "Sheet", 0, value="alpha"))
        assert not json.loads(M.excel_find_rows(session_key, "Sheet", 0, value="temporary"))
    finally:
        M.excel_close(session_key)
    assert hashlib.sha256(source.read_bytes()).hexdigest() == original_hash


@pytest.mark.parametrize("filename", [
    "01-audit-87-source.xlsx", "02-table-metadata-source.xlsx", "03-print-area-source.xlsx",
    "04-rich-text-phonetic-source.xlsx", "06-book1-richtext-source.xlsm", "07-real-package-source.xlsx",
])
def test_real_fixture_row_height_preserves_all_other_package_parts(tmp_path, filename):
    source = Path(r"D:\data-test\excel-preservation-fixtures\sources") / filename
    if not source.is_file():
        pytest.skip("external audit fixture unavailable")
    local = tmp_path / filename
    output = tmp_path / ("saved-" + filename)
    shutil.copyfile(source, local)
    original = local.read_bytes()
    key = _session_key(M.excel_load(str(local)))
    sheet = M._get_session(key)["sheets"][0]["name"]
    M.excel_set_dimension(key, sheet, "row", 0, 27)
    report = json.loads(M.excel_save(key, output_path=str(output), verify_preservation=False, report_format="json"))
    assert report["verification"]["preservation_ok"] is True
    assert inspect_xlsx_package(str(output))["valid"]
    with zipfile.ZipFile(local) as before, zipfile.ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert changed == ["xl/worksheets/sheet1.xml"]
    assert local.read_bytes() == original


def test_signed_fixture_row_height_requires_resigning(tmp_path):
    source = Path(r"D:\data-test\excel-preservation-fixtures\sources\05-advanced-package-source.xlsm")
    if not source.is_file():
        pytest.skip("external audit fixture unavailable")
    local = tmp_path / source.name
    output = tmp_path / ("saved-" + source.name)
    shutil.copyfile(source, local)
    original = local.read_bytes()
    key = _session_key(M.excel_load(str(local)))
    sheet = M._get_session(key)["sheets"][0]["name"]
    M.excel_set_dimension(key, sheet, "row", 0, 27)
    with pytest.raises(Exception, match="EXCEL_SAVE_REQUIRES_RESIGNING"):
        M.excel_save(key, output_path=str(output), verify_preservation=False)
    assert not output.exists()
    assert local.read_bytes() == original


def test_height_edits_on_two_sheets_leave_other_xml_unchanged(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "First"
    workbook.active["A1"] = "one"
    second = workbook.create_sheet("Second")
    second["A1"] = "two"
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    M.excel_set_dimension(key, "First", "row", 0, 24)
    M.excel_set_row_height(key, "Second", {"0": 30})
    report = json.loads(M.excel_save_as_copy(key, str(output), report_format="json", verify_preservation=True))
    assert report["verification"]["preservation_ok"] is True
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        changed = {name for name in before.namelist() if before.read(name) != after.read(name)}
        assert changed == {"xl/worksheets/sheet1.xml", "xl/worksheets/sheet2.xml"}


def test_reset_row_height_keeps_other_row_attributes(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active["A1"] = "value"
    workbook.active.row_dimensions[1].height = 28
    workbook.active.row_dimensions[1].hidden = True
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    M.excel_set_dimension(key, "S", "row", 0, None)
    report = json.loads(M.excel_save_as_copy(key, str(output), report_format="json", verify_preservation=True))
    assert report["verification"]["preservation_ok"] is True
    with zipfile.ZipFile(output) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
        assert b'hidden="1"' in sheet
        assert b" customHeight=" not in sheet
        assert b" ht=" not in sheet


@pytest.mark.parametrize("operation", ["sheet", "hyperlink", "table"])
def test_verified_package_additions_allow_only_expected_relationships(tmp_path, operation):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active["A1"] = "Name"
    workbook.active["B1"] = "Value"
    workbook.active["A2"] = "A"
    workbook.active["B2"] = 1
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    if operation == "sheet":
        M.excel_add_sheet(key, "Added")
    elif operation == "hyperlink":
        M.excel_set_hyperlink(key, "S", "A1", target="https://example.com")
    else:
        M.excel_add_table(key, "S", "AuditTable", "A1:B2")
    report = json.loads(M.excel_save_as_copy(key, str(output), report_format="json", verify_preservation=True))
    assert report["verification"]["preservation_ok"] is True
    assert inspect_xlsx_package(str(output))["valid"] is True


@pytest.mark.parametrize("relationship_id,relationship_type,old_content_type", [
    ("rId1", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet", False),
    ("rId2", "https://untrusted.example/relationships/worksheet", False),
    ("rId2", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet", True),
])
def test_package_additions_do_not_approve_modified_or_untrusted_relationships(
    relationship_id, relationship_type, old_content_type,
):
    base = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet"
    old_rel = {"type": base, "target": "worksheets/sheet1.xml", "target_mode": None}
    new_rel = {"type": relationship_type, "target": "worksheets/sheet2.xml", "target_mode": None}
    content_key = "Override:/xl/worksheets/sheet2.xml"
    content = {"ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"}
    before = SimpleNamespace(
        ensure_snapshot=lambda: {"workbook": {"sheets": [{"name": "S", "target": "xl/worksheets/sheet1.xml"}]}},
        parts={"xl/worksheets/sheet1.xml": b""},
        relationship_records={"xl/_rels/workbook.xml.rels#rId1": old_rel},
        content_type_records={content_key: content} if old_content_type else {},
    )
    after = SimpleNamespace(
        ensure_snapshot=lambda: {"workbook": {"sheets": [
            {"name": "S", "target": "xl/worksheets/sheet1.xml"},
            {"name": "Added", "target": "xl/worksheets/sheet2.xml"},
        ]}},
        parts={"xl/worksheets/sheet1.xml": b"", "xl/worksheets/sheet2.xml": b""},
        relationship_records={"xl/_rels/workbook.xml.rels#" + relationship_id: new_rel},
        content_type_records={content_key: content},
    )
    assert _requested_package_additions({"_dirty_paths": ["sheets/Added"], "sheets": []}, before, after) == []


def test_package_additions_do_not_approve_unrelated_worksheet_relationship():
    hyperlink_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
    style_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles"
    source = "xl/worksheets/_rels/sheet1.xml.rels"
    hyperlink = {"type": hyperlink_type, "target": "https://example.com", "target_mode": "External"}
    extra = {"type": style_type, "target": "../styles.xml", "target_mode": None}
    before = SimpleNamespace(
        ensure_snapshot=lambda: {"workbook": {"sheets": [{"name": "S", "target": "xl/worksheets/sheet1.xml"}]},
                                 "worksheets": {"S": {"relationships": []}}},
        parts={}, relationship_records={}, content_type_records={},
    )
    after = SimpleNamespace(
        ensure_snapshot=lambda: {"workbook": {"sheets": [{"name": "S", "target": "xl/worksheets/sheet1.xml"}]},
                                 "worksheets": {"S": {"relationships": [
                                     {"type": hyperlink_type, "target": "https://example.com", "external": True},
                                     {"type": style_type, "target": "xl/styles.xml", "external": False},
                                 ]}}},
        parts={}, relationship_records={f"{source}#rId1": hyperlink, f"{source}#rId2": extra},
        content_type_records={},
    )
    approved = _requested_package_additions({
        "_dirty_paths": ["sheets/S/hyperlinks/A1"],
        "sheets": [{"name": "S", "hyperlinks": {"A1": {"target": "https://example.com"}}}],
    }, before, after)
    assert approved == [f"package/relationships/{source}#rId1"]


def test_verified_multiple_hyperlinks_allow_only_their_relationships(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    M.excel_set_hyperlink(key, "S", "A1", target="https://example.com/one")
    M.excel_set_hyperlink(key, "S", "A2", target="https://example.com/two")
    report = json.loads(M.excel_save_as_copy(key, str(output), report_format="json", verify_preservation=True))
    assert report["verification"]["preservation_ok"] is True


def test_row_height_approval_does_not_hide_other_attributes_on_same_row(tmp_path):
    source = tmp_path / "source.xlsx"
    output = tmp_path / "output.xlsx"
    modified = tmp_path / "modified.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "S"
    workbook.active["A1"] = "value"
    workbook.save(source)
    key = _session_key(M.excel_load(str(source)))
    M.excel_set_dimension(key, "S", "row", 0, 27)
    report = json.loads(M.excel_save_as_copy(key, str(output), report_format="json", verify_preservation=True))
    assert report["verification"]["preservation_ok"] is True
    with zipfile.ZipFile(output) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml")
    assert b'customHeight="1"' in sheet
    _rewrite_part(output, modified, "xl/worksheets/sheet1.xml", b'customHeight="1"', b'customHeight="1" hidden="1"')
    verification = json.loads(M.excel_verify_preservation(
        str(modified), before_path=str(source),
        requested_paths=report["verification"]["requested_paths"],
    ))
    assert verification["preservation_ok"] is False
    assert any(item["path"] == "worksheets/S/rows/1/hidden" and item["classification"] == "UNAPPROVED_LOSS"
               for item in verification["changes"])
