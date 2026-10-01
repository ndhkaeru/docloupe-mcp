"""Exercise packaged MCP servers without host Python packages or optional CLIs."""

import json
import os
import shutil
import zipfile
from pathlib import Path

import anyio
import openpyxl
import pytest
from PIL import Image
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client


def _minimal_env():
    system_root = os.environ.get("SystemRoot", r"C:\Windows")
    env = {key: value for key, value in os.environ.items()
           if key.upper() not in {"PATH", "PYTHONHOME", "PYTHONPATH", "VIRTUAL_ENV", "NODE_PATH", "JAVA_HOME"}}
    env["PATH"] = os.path.join(system_root, "System32")
    env["PYTHONNOUSERSITE"] = "1"
    return env


def _binary(name):
    directory = os.environ.get("DOCLOUPE_MINIMAL_BIN_DIR")
    if not directory:
        pytest.skip("set DOCLOUPE_MINIMAL_BIN_DIR to test isolated binaries")
    binary = Path(directory) / (name + (".exe" if os.name == "nt" else ""))
    assert binary.is_file(), binary
    return str(binary.resolve())


def _docx(path):
    types = ('<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
             '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
             '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
             '</Types>')
    relationships = ('<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                     '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                     '</Relationships>')
    document = ('<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body><w:p><w:r><w:t>IsolatedDocxMarker</w:t></w:r></w:p></w:body></w:document>')
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document)


def _pdf(path):
    stream = b"BT /F1 12 Tf 50 700 Td (IsolatedPdfMarker) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for index, content in enumerate(objects, start=1):
        offsets.append(len(data))
        data += str(index).encode() + b" 0 obj\n" + content + b"\nendobj\n"
    start = len(data)
    data += b"xref\n0 6\n0000000000 65535 f \n"
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += b"trailer\n<< /Root 1 0 R /Size 6 >>\nstartxref\n" + str(start).encode() + b"\n%%EOF\n"
    path.write_bytes(data)


def _conversion_fixture(tmp_path, name):
    extension = {"csv-tools": ".csv", "docx-tools": ".docx", "html-tools": ".html",
                 "json-tools": ".json", "pdf-tools": ".pdf", "pptx-tools": ".pptx",
                 "text-tools": ".txt"}[name]
    path = tmp_path / (name + extension)
    marker = "Isolated" + name.split("-")[0].capitalize() + "Marker"
    if name == "docx-tools":
        _docx(path)
    elif name == "pdf-tools":
        _pdf(path)
    elif name == "pptx-tools":
        from pptx import Presentation
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = marker
        presentation.save(path)
    elif name == "csv-tools":
        path.write_text("Name\n" + marker + "\n", encoding="utf-8")
    elif name == "json-tools":
        path.write_text(json.dumps({"name": marker}), encoding="utf-8")
    elif name == "html-tools":
        path.write_text(f"<h1>{marker}</h1>", encoding="utf-8")
    else:
        path.write_text(marker, encoding="utf-8")
    return path, marker


@pytest.mark.parametrize("name", ["csv-tools", "docx-tools", "html-tools", "json-tools",
                                  "pdf-tools", "pptx-tools", "text-tools"])
def test_conversion_binary_without_host_dependencies(tmp_path, name):
    binary = _binary(name)
    path, marker = _conversion_fixture(tmp_path, name)

    async def check():
        with anyio.fail_after(60):
            async with stdio_client(StdioServerParameters(command=binary, env=_minimal_env())) as streams:
                async with ClientSession(*streams) as session:
                    await session.initialize()
                    assert {tool.name for tool in (await session.list_tools()).tools} == {"convert_to_markdown"}
                    result = await session.call_tool("convert_to_markdown", {"file_path": str(path)})
                    assert not result.isError, result.content
                    assert marker in result.content[0].text

    anyio.run(check)


def test_markdown_binary_without_host_dependencies(tmp_path):
    binary = _binary("md-tools")
    document = tmp_path / "diagram.md"
    document.write_text("# IsolatedMdMarker\n\n```mermaid\nflowchart TD\n A --> B\n```\n", encoding="utf-8")
    output = tmp_path / "diagram.svg"

    async def check():
        with anyio.fail_after(60):
            async with stdio_client(StdioServerParameters(command=binary, env=_minimal_env())) as streams:
                async with ClientSession(*streams) as session:
                    await session.initialize()
                    tools = {tool.name for tool in (await session.list_tools()).tools}
                    assert len(tools) == 45
                    assert "md_render_diagram" in tools
                    capabilities = await session.call_tool("md_runtime_capabilities")
                    assert capabilities.structuredContent["diagram"]["mermaid"]["backend"] == "rust"
                    assert capabilities.structuredContent["diagram"]["plantuml"]["validate"] is False
                    outline = await session.call_tool("markdown_outline", {"path": str(document)})
                    assert not outline.isError and "IsolatedMdMarker" in str(outline.content)
                    validation = await session.call_tool("md_validate_diagram", {"path": str(document)})
                    assert not validation.isError and validation.structuredContent["ok"] is True
                    render = await session.call_tool("md_render_diagram", {"path": str(document), "output_path": str(output)})
                    assert not render.isError and render.structuredContent["ok"] is True
                    assert "<svg" in output.read_text(encoding="utf-8")

    anyio.run(check)


def test_every_markdown_tool_without_host_dependencies(tmp_path):
    binary = _binary("md-tools")
    document_text = (
        "---\ntitle: Fixture\n---\n# Main\n\nIntro [link](other.md) ![image](image.png).\n\n"
        "## Child\n\n| Name | Value |\n| --- | --- |\n| First | 1 |\n\n"
        "```mermaid\nflowchart TD\n A --> B\n```\n\n"
        "```python\nprint('fixture')\n```\n\n## Next\n\nEnd.\n"
    )
    (tmp_path / "other.md").write_text("# Other\n", encoding="utf-8")
    (tmp_path / "image.png").write_bytes(b"image")

    def arguments(name, path):
        common = {"path": str(path)}
        overrides = {
            "read_markdown_section": {"heading": "Child"},
            "md_read_range": {"start_line": 1, "end_line": 3},
            "md_read_near": {"text": "Intro"},
            "replace_markdown_section": {"heading": "Child", "new_content": "Updated.\n"},
            "md_search": {"query": "Intro"},
            "md_insert_section": {"title": "Added", "content": "Content.\n"},
            "md_delete_section": {"heading": "Child"},
            "md_append_to_section": {"heading": "Child", "content": "Appended.\n"},
            "md_rename_heading": {"heading": "Child", "new_title": "Renamed"},
            "md_replace_text": {"find": "Intro", "replace": "Replaced"},
            "md_edit_table": {"op": "set_cell", "row": 0, "col": 0, "value": "Second"},
            "md_insert_table": {"headers": ["Column"], "rows": [["Value"]]},
            "md_insert_diagram": {"source": "flowchart TD\n B --> C"},
            "md_replace_diagram": {"source": "flowchart TD\n A --> C"},
            "md_replace_code_block": {"source": "flowchart TD\n A --> C"},
            "md_insert_code_block": {"source": "print('another')", "info": "python"},
            "md_get_anchor": {"title": "Child"},
            "md_patch_lines": {"start_line": 6, "end_line": 6, "replacement": "Replaced intro.\n"},
            "md_set_heading_level": {"heading": "Child", "level": 3},
            "md_move_section": {"heading": "Child", "target_heading": "Next"},
            "md_rewrite_links": {"replacements": {"other.md": "renamed.md"}},
            "md_frontmatter": {"op": "read"},
            "md_tangle": {"output_dir": str(tmp_path / "tangle")},
            "md_split": {"output_dir": str(tmp_path / "split")},
            "md_merge": {"paths": [str(path)], "output_path": str(tmp_path / "merged.md")},
            "md_render_diagram": {"output_path": str(tmp_path / "render.svg")},
        }
        if name in {"md_get_anchor", "md_merge", "md_runtime_capabilities"}:
            return overrides.get(name, {})
        return common | overrides.get(name, {})

    async def check():
        with anyio.fail_after(180):
            async with stdio_client(StdioServerParameters(command=binary, env=_minimal_env())) as streams:
                async with ClientSession(*streams) as session:
                    await session.initialize()
                    names = {tool.name for tool in (await session.list_tools()).tools}
                    assert len(names) == 45
                    failures = []
                    for name in sorted(names):
                        document = tmp_path / f"{name}.md"
                        document.write_text(document_text, encoding="utf-8")
                        result = await session.call_tool(name, arguments(name, document))
                        if result.isError:
                            failures.append((name, result.content))
                        if name == "md_render_diagram" and not result.isError:
                            assert result.structuredContent["ok"] is True
                            assert (tmp_path / "render.svg").is_file()
                    assert not failures, failures

    anyio.run(check)


def test_excel_binary_without_host_dependencies(tmp_path):
    binary = _binary("excel-tools")
    workbook_path = tmp_path / "source.xlsx"
    output = tmp_path / "result.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "IsolatedExcelMarker"
    workbook.save(workbook_path)

    async def check():
        with anyio.fail_after(90):
            async with stdio_client(StdioServerParameters(command=binary, env=_minimal_env())) as streams:
                async with ClientSession(*streams) as session:
                    await session.initialize()
                    tools = {tool.name for tool in (await session.list_tools()).tools}
                    assert len(tools) == 112
                    info = await session.call_tool("excel_get_info", {"uri": str(workbook_path)})
                    assert not info.isError and "Sheet" in str(info.content)
                    loaded = await session.call_tool("excel_load", {"uri": str(workbook_path)})
                    assert not loaded.isError, loaded.content
                    session_key = str(workbook_path.resolve())
                    cell = await session.call_tool("excel_get_cell", {
                        "session_key": session_key, "sheet_name": "Sheet", "row_index": 0, "col_index": 0,
                    })
                    assert not cell.isError and "IsolatedExcelMarker" in str(cell.content)
                    saved = await session.call_tool("excel_save_as_copy", {
                        "session_key": session_key, "output_path": str(output), "verify_preservation": True,
                    })
                    assert not saved.isError, saved.content
                    assert output.is_file()
                    closed = await session.call_tool("excel_close", {"session_key": session_key})
                    assert not closed.isError

    anyio.run(check)


def test_every_excel_tool_without_host_dependencies(tmp_path):
    binary = _binary("excel-tools")
    source = tmp_path / "source.xlsx"
    identical = tmp_path / "identical.xlsx"
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.append(["Label", "Amount"])
    worksheet.append(["IsolatedExcelMarker", 1])
    worksheet.append(["Second", 2])
    workbook.save(source)
    shutil.copyfile(source, identical)
    image = tmp_path / "image.png"
    Image.new("RGB", (2, 2), color="red").save(image)

    async def check():
        with anyio.fail_after(600):
            async with stdio_client(StdioServerParameters(command=binary, env=_minimal_env())) as streams:
                async with ClientSession(*streams) as session:
                    await session.initialize()
                    tools = {tool.name: tool for tool in (await session.list_tools()).tools}
                    assert len(tools) == 112
                    seen = set()
                    session_key = str(source.resolve())
                    defaults = {"session_key": session_key, "sheet_name": "Sheet"}

                    async def run(name, overrides=None, expected_error=None):
                        required = tools[name].inputSchema.get("required", [])
                        arguments = {field: defaults[field] for field in required if field in defaults}
                        arguments.update(overrides or {})
                        result = await session.call_tool(name, arguments)
                        seen.add(name)
                        if expected_error:
                            assert result.isError and expected_error in str(result.content), (name, result.content)
                        else:
                            assert not result.isError, (name, result.content)
                        return result

                    await run("convert_to_markdown", {"file_path": str(source)})
                    await run("excel_get_info", {"uri": str(source)})
                    await run("excel_get_workbook_summary", {"file_path": str(source)})
                    await run("excel_get_sheet_preview", {"file_path": str(source)})
                    await run("excel_validate_workbook", {"path": str(source)})
                    await run("excel_diff_package", {"before_path": str(source), "after_path": str(identical)})
                    await run("excel_verify_preservation", {"before_path": str(source), "after_path": str(identical)})
                    await run("excel_build_preservation_summary", {
                        "coverage_reports": [{"coverage_id": "minimal", "checked_count": 1, "passed_count": 1}],
                        "verification_reports": [{"fixture_id": "minimal", "preservation_ok": True,
                                                  "fixed_output_preserved": True, "fixture_graph_valid": True}],
                        "backup_checks": [{"check_id": "minimal", "passed": True}],
                    })
                    await run("excel_load", {"uri": str(source)})

                    reads = [
                        ("excel_get_session_status", {}), ("excel_to_markdown", {}),
                        ("excel_to_markdown_range", {"range_ref": "A1:B3"}),
                        ("excel_get_rows", {}), ("excel_read_range", {"range_ref": "A1:B3"}),
                        ("excel_find_cells", {"query": "IsolatedExcelMarker"}),
                        ("excel_list_tables", {}), ("excel_list_defined_names", {}),
                        ("excel_get_cell", {"row_index": 1, "col_index": 0}),
                        ("excel_get_column", {"col_index": 0}),
                        ("excel_get_shapes", {}), ("excel_get_rich_text", {"cell": "A2"}),
                        ("excel_get_conditional_formats", {}),
                        ("excel_find_rows", {"col_index": 0, "value": "Second"}),
                        ("excel_get_workbook_semantics", {}), ("excel_get_workbook_views", {}),
                        ("excel_get_sheet_semantics", {}), ("excel_get_sheet_views", {}),
                        ("excel_list_package_parts", {}),
                        ("excel_read_package_part", {"part_path": "xl/workbook.xml"}),
                        ("excel_extract_images", {"output_dir": str(tmp_path / "images")}),
                    ]
                    for name, arguments in reads:
                        await run(name, arguments)

                    await run("excel_edit_cells", {"edits": [{"cell": "C2", "value": "edited"}]})
                    saved_copy = tmp_path / "saved.xlsx"
                    await run("excel_save_as_copy", {"output_path": str(saved_copy),
                                                     "verify_preservation": True})
                    assert saved_copy.is_file()
                    await run("excel_save", {"verify_preservation": True})
                    await run("excel_reload")
                    await run("excel_capture", {"output_path": str(tmp_path / "capture.png"),
                                                "soffice_path": str(tmp_path / "missing-soffice.exe")},
                              expected_error="[WinError 2]")

                    edits = [
                        ("excel_set_style", {"r1": 1, "c1": 0, "style": {"bold": True}}),
                        ("excel_set_font_color", {"r1": 1, "c1": 0, "color": "FFFF0000"}),
                        ("excel_set_strike", {"r1": 1, "c1": 0, "enabled": True}),
                        ("excel_set_borders", {"r1": 1, "c1": 0, "r2": 1, "c2": 0,
                                               "sides": ["bottom"], "style": "thin"}),
                        ("excel_set_dimension", {"axis": "row", "index": 1, "size": 24}),
                        ("excel_set_row_height", {"row_heights": {"1": 24}}),
                        ("excel_set_column_width", {"col_widths": {"A": 18}}),
                        ("excel_autofit_cols", {"col_indices": [0]}),
                        ("excel_freeze_panes", {"row": 1, "col": 0}),
                        ("excel_set_data_validation", {"start_row": 1, "end_row": 2,
                                                       "start_col": 0, "end_col": 0,
                                                       "options": ["One", "Two"]}),
                        ("excel_fill_column", {"col_index": 5, "start_row": 1,
                                               "end_row": 2, "value": "filled"}),
                        ("excel_set_formula", {"cell": "D2", "formula": "=B2+1"}),
                        ("excel_edit_rich_text", {"cell": "G2", "operations": [
                            {"op": "replace_runs", "runs": [{"text": "rich", "font": {}}]},
                        ]}),
                        ("excel_set_auto_filter", {"ref": "A1:B3"}),
                        ("excel_set_ignored_errors", {"rules": []}),
                        ("excel_set_cell_style_semantics", {"range_ref": "A2", "xf": {"applyFont": True}}),
                        ("excel_set_calculation_properties", {"properties": {"fullCalcOnLoad": True}}),
                        ("excel_set_workbook_properties", {"properties": {"date1904": False}}),
                        ("excel_set_document_properties", {"core": {"title": "Minimal"}}),
                        ("excel_set_workbook_protection", {"properties": {"lockStructure": False}}),
                        ("excel_set_workbook_views", {"views": [{"activeTab": 0}]}),
                        ("excel_set_sheet_properties", {"properties": {"codeName": "Sheet"}}),
                        ("excel_set_sheet_views", {"views": [{"showGridLines": True}]}),
                        ("excel_set_row_properties", {"row_index": 1, "properties": {"height": 24}}),
                        ("excel_set_page_setup", {"properties": {"orientation": "landscape"}}),
                        ("excel_set_print_options", {"properties": {"gridLines": False}}),
                        ("excel_set_header_footer", {"sections": {"oddHeader": "Header"}}),
                        ("excel_set_page_breaks", {"row_breaks": []}),
                        ("excel_set_print_area", {"areas": "A1:B3"}),
                        ("excel_set_print_titles", {"repeated_rows": "1:1"}),
                        ("excel_set_sheet_protection", {"enabled": False}),
                        ("excel_set_protected_ranges", {"ranges": []}),
                        ("excel_set_hyperlink", {"cell": "H2", "target": "https://example.com"}),
                        ("excel_remove_hyperlink", {"cell": "H2"}),
                        ("excel_set_comment", {"cell": "H2", "text": "note"}),
                        ("excel_remove_comment", {"cell": "H2"}),
                    ]
                    for name, arguments in edits:
                        await run(name, arguments)

                    await run("excel_add_defined_name", {"name": "MinimalName", "value": "Sheet!$A$2"})
                    await run("excel_update_defined_name", {"name": "MinimalName", "updates": {"value": "Sheet!$B$2"}})
                    await run("excel_delete_defined_name", {"name": "MinimalName"})
                    await run("excel_add_named_style", {"name": "MinimalStyle", "style": {"bold": True}})
                    await run("excel_update_named_style", {"name": "MinimalStyle", "updates": {"italic": True}})
                    await run("excel_delete_named_style", {"name": "MinimalStyle"})
                    await run("excel_add_table", {"name": "MinimalTable", "ref": "A1:B3"})
                    await run("excel_update_table", {"name": "MinimalTable", "updates": {"ref": "A1:B3"}})
                    await run("excel_delete_table", {"name": "MinimalTable"})
                    result = await run("excel_add_conditional_format", {"sqref": "B2:B3", "rule": {
                        "type": "expression", "formula": ["B2>0"],
                    }})
                    rule_id = json.loads(result.content[0].text)["rule_id"]
                    await run("excel_update_conditional_format", {"rule_id": rule_id,
                                                                  "updates": {"formula": ["B2>1"]}})
                    await run("excel_delete_conditional_format", {"rule_id": rule_id})

                    await run("excel_add_sheet", {"sheet_name": "Extra"})
                    await run("excel_rename_sheet", {"sheet_name": "Extra", "new_name": "Renamed"})
                    await run("excel_move_sheet", {"sheet_name": "Renamed", "position": 0})
                    await run("excel_delete_sheet", {"sheet_name": "Renamed"})
                    await run("excel_copy_sheet", {"source_sheet": "Sheet", "new_name": "Copied"})
                    await run("excel_delete_sheet", {"sheet_name": "Copied"})
                    destination = tmp_path / "destination.xlsx"
                    shutil.copyfile(source, destination)
                    await run("excel_load", {"uri": str(destination)})
                    await run("excel_copy_sheet_to", {
                        "src_session_key": session_key, "src_sheet_name": "Sheet",
                        "dst_session_key": str(destination.resolve()), "new_name": "Transferred",
                    })
                    await run("excel_close", {"session_key": str(destination.resolve())})

                    cloned = await run("excel_clone_rows", {"start_row": 1})
                    await run("excel_insert_rows", {"inserts": [{
                        "after_index": 2, "rows_json": json.loads(cloned.content[0].text),
                    }]})
                    await run("excel_copy_row", {"row_index": 1, "after_index": 2})
                    await run("excel_delete_rows", {"row_indices": [3]})
                    await run("excel_fill_rows", {"template_row": 1, "after_index": 2, "count": 1})
                    await run("excel_insert_column", {"after_col_index": 3})
                    await run("excel_copy_column", {"col_index": 0, "after_col_index": 2})
                    await run("excel_delete_column", {"col_index": 3})
                    await run("excel_clear_range", {"r1": 5, "c1": 5, "r2": 5, "c2": 5})
                    await run("excel_merge_cells", {"r1": 5, "c1": 5, "r2": 5, "c2": 6})
                    await run("excel_set_sheet_state", {"state": "visible"})

                    await run("excel_upsert_package_part", {"part_path": "customXml/minimal.xml",
                                                            "content": "<minimal/>", "content_type": "application/xml"})
                    await run("excel_set_package_relationships", {"source_part": "customXml/minimal.xml",
                                                                  "relationships": []})
                    await run("excel_set_package_content_types", {"defaults": [
                        {"extension": "txt", "content_type": "text/plain"},
                    ]})
                    await run("excel_apply_package_transaction", {"upsert": [{
                        "path": "customXml/transaction.xml", "content": "<transaction/>",
                        "content_type": "application/xml",
                    }]})
                    await run("excel_delete_package_part", {"part_path": "customXml/minimal.xml"})

                    created = await run("excel_create_workbook", {
                        "sheet_names": ["Drawing"], "target_path": str(tmp_path / "drawing.xlsx"),
                    })
                    drawing_key = json.loads(created.content[0].text)["session_key"]
                    draw = {"session_key": drawing_key, "sheet_name": "Drawing"}
                    await run("excel_add_shape", draw | {"shape_type": "rect", "anchor": "D2", "text": "Before"})
                    await run("excel_add_image", draw | {"anchor": "D8", "source_path": str(image)})
                    await run("excel_add_chart", draw | {"chart_type": "bar", "source_range": "A1:B3",
                                                         "anchor": "H2"})
                    await run("excel_save", {"session_key": drawing_key})
                    await run("excel_close", {"session_key": drawing_key})
                    drawing_path = tmp_path / "drawing.xlsx"
                    await run("excel_load", {"uri": str(drawing_path)})
                    loaded_drawing = {"session_key": str(drawing_path.resolve()), "sheet_name": "Drawing"}
                    shapes = await run("excel_get_shapes", loaded_drawing)
                    shape_index = next(shape["index"] for shape in json.loads(shapes.content[0].text)["Drawing"]
                                       if shape["type"] == "shape")
                    await run("excel_update_shape_text", loaded_drawing | {"shape_index": shape_index, "text": "After"})
                    await run("excel_set_shape_style", loaded_drawing | {"shape_index": shape_index,
                                                                         "fill_color": "FF00FF00"})
                    await run("excel_close", {"session_key": str(drawing_path.resolve())})
                    await run("excel_close")
                    assert seen == set(tools), sorted(set(tools) - seen)

    anyio.run(check)
